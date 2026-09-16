import os
import re
from typing import Dict, Any

try:
    from pypdf import PdfReader
except ImportError:
    try:
        from PyPDF2 import PdfReader
    except ImportError:
        PdfReader = None

try:
    import docx
except ImportError:
    docx = None

class ResumeParser:
    """
    Extracts raw text and basic metadata from PDF, DOCX, and TXT resumes.
    Includes text cleaning and sanitization.
    """

    ALLOWED_EXTENSIONS = {'.pdf', '.docx', '.txt'}

    @staticmethod
    def extract_text(file_path: str) -> str:
        """
        Main entry point to extract text from a file path based on extension.
        """
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"Resume file not found at: {file_path}")

        ext = os.path.splitext(file_path)[1].lower()

        if ext == '.pdf':
            return ResumeParser._extract_from_pdf(file_path)
        elif ext == '.docx':
            return ResumeParser._extract_from_docx(file_path)
        elif ext == '.txt':
            return ResumeParser._extract_from_txt(file_path)
        else:
            raise ValueError(f"Unsupported file format: '{ext}'. Allowed formats: PDF, DOCX, TXT.")

    @staticmethod
    def _extract_from_pdf(file_path: str) -> str:
        if PdfReader is None:
            raise ImportError("pypdf or PyPDF2 is required for reading PDF files. Please install pypdf.")
        
        text_content = []
        try:
            reader = PdfReader(file_path)
            for page_idx, page in enumerate(reader.pages):
                extracted = page.extract_text()
                if extracted:
                    text_content.append(extracted)
        except Exception as e:
            raise ValueError(f"Failed to read PDF file: {str(e)}")

        full_text = "\n".join(text_content).strip()
        if not full_text:
            raise ValueError("The uploaded PDF file contains no selectable text or is scanned/empty.")
        
        return ResumeParser.clean_text(full_text)

    @staticmethod
    def _extract_from_docx(file_path: str) -> str:
        if docx is None:
            raise ImportError("python-docx is required for reading DOCX files. Please install python-docx.")
        
        text_content = []
        try:
            doc = docx.Document(file_path)
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text_content.append(paragraph.text)
            for table in doc.tables:
                for row in table.rows:
                    row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_text:
                        text_content.append(" | ".join(row_text))
        except Exception as e:
            raise ValueError(f"Failed to read DOCX file: {str(e)}")

        full_text = "\n".join(text_content).strip()
        if not full_text:
            raise ValueError("The uploaded DOCX file is empty or contains no extractable text.")

        return ResumeParser.clean_text(full_text)

    @staticmethod
    def _extract_from_txt(file_path: str) -> str:
        try:
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                content = f.read().strip()
        except Exception as e:
            raise ValueError(f"Failed to read text file: {str(e)}")

        if not content:
            raise ValueError("The uploaded text file is empty.")

        return ResumeParser.clean_text(content)

    @staticmethod
    def clean_text(text: str) -> str:
        """
        Cleans and normalizes raw extracted text.
        """
        if not text:
            return ""
        # Remove null bytes and unusual control characters
        text = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', text)
        # Normalize excessive vertical whitespace
        text = re.sub(r'\n{3,}', '\n\n', text)
        # Replace multiple horizontal spaces/tabs
        text = re.sub(r'[ \t]+', ' ', text)
        return text.strip()

    @staticmethod
    def extract_metadata(text: str) -> Dict[str, Any]:
        """
        Extracts quick structural stats (word count, email/phone heuristics if present).
        """
        words = text.split()
        emails = re.findall(r'[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}', text)
        phones = re.findall(r'\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}', text)
        
        return {
            "word_count": len(words),
            "char_count": len(text),
            "detected_emails": list(set(emails)),
            "detected_phones": list(set(phones))
        }
